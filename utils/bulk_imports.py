from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from zipfile import BadZipFile


def excel_response(filename, worksheet_name, headers, rows):
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=worksheet_name)
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=header)
        cell.font = Font(bold=True)
        header_cells.append(cell)
    worksheet.append(header_cells)

    for row in rows:
        cells = []
        for value in row:
            cell = WriteOnlyCell(worksheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        worksheet.append(cells)

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def get_request_school(request):
    tenant = getattr(request, "tenant", None)
    school = getattr(tenant, "school", None)
    if school is not None:
        return school

    user = getattr(request, "user", None)
    if user is not None and user.is_authenticated:
        return getattr(user, "school", None)
    return None


def iter_excel_rows(upload, columns):
    try:
        workbook = openpyxl.load_workbook(upload, read_only=True, data_only=True)
    except (BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError("Upload a valid .xlsx workbook.") from exc
    try:
        worksheet = workbook.active
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not any(value not in (None, "") for value in values):
                continue
            padded_values = list(values[: len(columns)])
            padded_values.extend([None] * (len(columns) - len(padded_values)))
            yield row_number, dict(zip(columns, padded_values))
    finally:
        workbook.close()


def text_value(value):
    if value is None:
        return ""
    return str(value).strip()


def date_value(value, field_name):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_date(text_value(value))
    if parsed is None:
        raise ValueError(f"{field_name} must be a valid date in YYYY-MM-DD format.")
    return parsed


def decimal_value(value, field_name):
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, AttributeError):
        raise ValueError(f"{field_name} must be a valid number.") from None
